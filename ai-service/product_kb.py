import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

import requests


ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "Mapa de Produtos __ Marketing.xlsx"
SITE_URL = "https://www.silicon.ind.br/"

# Cache simples para evitar releitura do Excel em cada chamada
_EXCEL_CACHE: Dict[str, Any] = {
    "ok": False,
    "items": [],
    "sheets": [],
    "error": None,
    "mtime": None,
}

def _get_excel_mtime() -> Optional[float]:
    try:
        return EXCEL_PATH.stat().st_mtime
    except Exception:
        return None


def _safe_lower(s: Optional[str]) -> str:
    return (s or "").strip().lower()


def _try_import_openpyxl():
    try:
        import openpyxl  # type: ignore
        return openpyxl
    except Exception:
        return None


def _strip_html(html: str) -> str:
    # remove scripts/styles
    html = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.IGNORECASE)
    html = re.sub(r"<style[\s\S]*?</style>", " ", html, flags=re.IGNORECASE)
    # remove tags
    text = re.sub(r"<[^>]+>", " ", html)
    # normalize spaces
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _detect_category(text: str) -> str:
    t = _safe_lower(text)
    if any(k in t for k in ["solar", "fotovolta", "usina", "energia"]):
        return "solar"
    if any(k in t for k in ["led", "ilumina", "luminária", "luminaria", "projetos luminotécnicos"]):
        return "led"
    return "geral"


def load_excel_index() -> Dict[str, Any]:
    """Carrega TODAS as abas do Excel em uma estrutura simples de busca.

    Retorna um dicionário com:
      - ok: bool
      - items: lista de registros estruturados
      - sheets: nomes das abas lidas
      - error: mensagem se falhar
    """
    items: List[Dict[str, Any]] = []
    sheets: List[str] = []

    # Verifica existência e mtime para cache
    if not EXCEL_PATH.exists():
        return {"ok": False, "items": [], "sheets": sheets, "error": f"Arquivo não encontrado: {EXCEL_PATH}"}

    current_mtime = _get_excel_mtime()
    if (
        _EXCEL_CACHE.get("ok")
        and _EXCEL_CACHE.get("mtime") == current_mtime
        and isinstance(_EXCEL_CACHE.get("items"), list)
    ):
        # Retorna cache quente
        return {
            "ok": True,
            "items": _EXCEL_CACHE.get("items", []),
            "sheets": _EXCEL_CACHE.get("sheets", []),
        }

    ox = _try_import_openpyxl()
    if not ox:
        return {"ok": False, "items": [], "sheets": sheets, "error": "Pacote openpyxl não instalado."}

    try:
        wb = ox.load_workbook(EXCEL_PATH)
        for ws in wb.worksheets:
            sheets.append(ws.title)
            # coletar cabeçalhos da primeira linha
            headers: List[str] = []
            if ws.max_row < 2:
                continue
            for j, cell in enumerate(ws[1], start=1):
                headers.append(_safe_lower(str(cell.value)))

            def get_field(row_dict: Dict[str, Any], candidates: List[str]) -> Optional[str]:
                for c in candidates:
                    if c in row_dict and row_dict[c]:
                        return str(row_dict[c])
                return None

            for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                row_dict: Dict[str, Any] = {}
                for j, cell in enumerate(row):
                    key = headers[j] if j < len(headers) else f"col_{j}"
                    row_dict[key] = cell.value

                name = get_field(row_dict, ["nome", "produto", "linha", "modelo"]) or ""
                desc = get_field(row_dict, ["descrição", "descricao", "resumo", "aplicação", "aplicacao"]) or ""
                difs = get_field(row_dict, ["diferenciais", "características", "caracteristicas"]) or ""
                segm = get_field(row_dict, ["segmento", "setor"]) or ""
                link = get_field(row_dict, ["link", "url", "site"]) or ""

                text_blob = " ".join([str(name), str(desc), str(difs), str(segm)]).strip()
                item = {
                    "name": str(name).strip(),
                    "description": str(desc).strip(),
                    "features": str(difs).strip(),
                    "segment": str(segm).strip(),
                    "link": str(link).strip(),
                    "category": _detect_category(text_blob),
                    "sheet": ws.title,
                    "_raw": row_dict,
                }
                if any(v for k, v in item.items() if k in ("name", "description", "features")):
                    items.append(item)

        # Tenta carregar planilhas adicionais na raiz
        extra_paths = [p for p in ROOT.glob("*.xlsx") if p.name != EXCEL_PATH.name and not p.name.startswith("~$")]
        for xp in extra_paths:
            try:
                wb2 = ox.load_workbook(xp)
                for ws in wb2.worksheets:
                    sheets.append(f"{xp.name}:{ws.title}")
                    headers: List[str] = []
                    if ws.max_row < 2:
                        continue
                    for j, cell in enumerate(ws[1], start=1):
                        headers.append(_safe_lower(str(cell.value)))
                    def get_field2(row_dict: Dict[str, Any], candidates: List[str]) -> Optional[str]:
                        for c in candidates:
                            if c in row_dict and row_dict[c]:
                                return str(row_dict[c])
                        return None
                    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                        row_dict: Dict[str, Any] = {}
                        for j, cell in enumerate(row):
                            key = headers[j] if j < len(headers) else f"col_{j}"
                            row_dict[key] = cell.value
                        name = get_field2(row_dict, ["nome", "produto", "linha", "modelo"]) or ""
                        desc = get_field2(row_dict, ["descrição", "descricao", "resumo", "aplicação", "aplicacao"]) or ""
                        difs = get_field2(row_dict, ["diferenciais", "características", "caracteristicas"]) or ""
                        segm = get_field2(row_dict, ["segmento", "setor"]) or ""
                        link = get_field2(row_dict, ["link", "url", "site"]) or ""
                        text_blob = " ".join([str(name), str(desc), str(difs), str(segm)]).strip()
                        item = {
                            "name": str(name).strip(),
                            "description": str(desc).strip(),
                            "features": str(difs).strip(),
                            "segment": str(segm).strip(),
                            "link": str(link).strip(),
                            "category": _detect_category(text_blob),
                            "sheet": f"{xp.name}:{ws.title}",
                            "_raw": row_dict,
                        }
                        if any(v for k, v in item.items() if k in ("name", "description", "features")):
                            items.append(item)
            except Exception:
                pass

        # Atualiza cache
        _EXCEL_CACHE["ok"] = True
        _EXCEL_CACHE["items"] = items
        _EXCEL_CACHE["sheets"] = sheets
        _EXCEL_CACHE["mtime"] = current_mtime

        return {"ok": True, "items": items, "sheets": sheets}
    except Exception as e:
        # Atualiza cache com erro (não sobrescreve itens válidos prévios)
        _EXCEL_CACHE["ok"] = False
        _EXCEL_CACHE["error"] = str(e)
        _EXCEL_CACHE["mtime"] = current_mtime
        return {"ok": False, "items": [], "sheets": sheets, "error": f"Falha ao ler Excel: {e}"}


def load_site_text() -> Dict[str, Any]:
    try:
        resp = requests.get(SITE_URL, timeout=10)
        resp.raise_for_status()
        text = _strip_html(resp.text)
        return {"ok": True, "text": text, "url": SITE_URL}
    except Exception as e:
        return {"ok": False, "text": "", "url": SITE_URL, "error": str(e)}


def _extract_links(html: str, base: str) -> List[str]:
    links: List[str] = []
    for m in re.finditer(r"href=\"([^\"]+)\"", html, flags=re.IGNORECASE):
        href = m.group(1).strip()
        if href.startswith("#") or href.startswith("mailto:") or href.startswith("tel:"):
            continue
        if href.startswith("http"):
            if "silicon.ind.br" not in href:
                continue
            url = href
        else:
            # relativo
            if not href.startswith("/"):
                href = "/" + href
            url = base.rstrip("/") + href
        # filtra extensões não-texto
        if re.search(r"\.(pdf|jpg|jpeg|png|gif|svg|zip|rar)$", url, flags=re.IGNORECASE):
            continue
        links.append(url.split("#")[0])
    # remove duplicados
    uniq: List[str] = []
    seen = set()
    for l in links:
        if l not in seen:
            seen.add(l)
            uniq.append(l)
    return uniq


def crawl_site(max_pages: int = 20) -> Dict[str, Any]:
    """Rastreia a home e até max_pages páginas internas do domínio Silicon.

    Retorna {ok, pages: [{url, text}], error?}
    """
    pages: List[Dict[str, str]] = []
    try:
        first = requests.get(SITE_URL, timeout=10)
        first.raise_for_status()
        home_html = first.text
        pages.append({"url": SITE_URL, "text": _strip_html(home_html)})
        to_visit = _extract_links(home_html, SITE_URL)[:max_pages]
        for url in to_visit:
            try:
                r = requests.get(url, timeout=10)
                if r.status_code == 200:
                    pages.append({"url": url, "text": _strip_html(r.text)})
            except Exception:
                pass
        return {"ok": True, "pages": pages}
    except Exception as e:
        return {"ok": False, "pages": pages, "error": str(e)}


def _score_item(query: str, item: Dict[str, Any]) -> float:
    q = _safe_lower(query)
    text = " ".join([
        _safe_lower(item.get("name")),
        _safe_lower(item.get("description")),
        _safe_lower(item.get("features")),
        _safe_lower(item.get("segment")),
        _safe_lower(item.get("category")),
    ])
    score = 0.0
    for token in re.split(r"\W+", q):
        if token and token in text:
            score += 1.0
    return score


def _search_site_snippets(query: str, text: str, top_k: int = 3) -> List[str]:
    q = _safe_lower(query)
    # Segmenta de forma ingênua em frases por ponto
    sentences = [s.strip() for s in re.split(r"[.!?]", text) if s.strip()]
    scored: List[tuple[float, str]] = []
    for s in sentences:
        s_low = _safe_lower(s)
        score = 0.0
        for token in re.split(r"\W+", q):
            if token and token in s_low:
                score += 1.0
        if score > 0:
            scored.append((score, s))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [s for _, s in scored[:top_k]]


def query_product_knowledge(query: str, top_k: int = 3) -> Dict[str, Any]:
    """Consulta conhecimento de produtos usando exclusivamente o Excel.

    Retorna dicionário com:
      - excel_loaded: bool
      - excel_results: lista de itens relevantes
      - hints: mensagens de utilidade (ex.: instalar openpyxl)
    """
    hints: List[str] = []

    excel = load_excel_index()
    excel_results: List[Dict[str, Any]] = []
    if excel.get("ok"):
        items = excel.get("items", [])
        scored = [( _score_item(query, it), it ) for it in items]
        scored = [p for p in scored if p[0] > 0]
        scored.sort(key=lambda x: x[0], reverse=True)
        excel_results = [it for _, it in scored[:top_k]]
    else:
        err = excel.get("error")
        if err:
            hints.append(str(err))

    return {
        "excel_loaded": bool(excel.get("ok")),
        "excel_results": excel_results,
        "hints": hints,
    }


if __name__ == "__main__":
    # teste rápido manual
    q = "luminária LED industrial Audace Pro"
    print(json.dumps(query_product_knowledge(q), ensure_ascii=False, indent=2))